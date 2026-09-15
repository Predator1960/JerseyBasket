"""
JerseyBasket product database sync tool
=========================================
Rebuilds the "Products" sheet of JerseyBasket-Product-Database.xlsx from a
fresh CSV export of the live app's BASE_PRODUCTS.

USAGE:
    python sync_products.py

Expects two files in the SAME FOLDER as this script:
    - product-full-export.csv   (get this from Claude Code — see prompt below)
    - JerseyBasket-Product-Database.xlsx

Overwrites the Products sheet in place; all other sheets (SPENDS, Summary,
review sheets) are left untouched. Recalculates Lowest Price / Cheapest
Store / App Last Updated automatically. Preserves per-item icons from the
existing file where a product name already exists; falls back to the
category's icon for new products.

TO GET A FRESH CSV, ask Claude Code:
    "Export the full BASE_PRODUCTS array from src/App.jsx as a CSV with
    these columns: ID, Product Name, Category, Co-op Price, Morrisons Price,
    M&S Price, Waitrose Price, Iceland Price, Alliance Price. Use blank for
    any store with no price. Save as product-full-export.csv in the
    project folder." Then copy that CSV into this script's folder.
"""

import openpyxl, csv, re, sys, os
from datetime import date
from openpyxl.styles import Font, PatternFill

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(SCRIPT_DIR, "product-full-export.csv")
XLSX_PATH = os.path.join(SCRIPT_DIR, "JerseyBasket-Product-Database.xlsx")

STORE_DISPLAY = ["CI Co-op", "Morrisons", "M&S Food", "Waitrose", "Iceland", "Alliance"]

# Price column -> that store's normal (non-cheapest) background fill.
# Matches the sheet's existing per-column zebra colours; the cheapest
# cell(s) in each row instead get CHEAPEST_FILL/CHEAPEST_FONT below.
STORE_COLS = {"CI Co-op": 5, "Morrisons": 7, "M&S Food": 9, "Waitrose": 11, "Iceland": 13, "Alliance": 15}
NORMAL_FILL = {
    "CI Co-op":  "FFDBEAFE",
    "Morrisons": "FFFEF9C3",
    "M&S Food":  "FFF1F5F9",
    "Waitrose":  "FFCCFBF1",
    "Iceland":   "FFFEE2E2",
    "Alliance":  "FFFFE4E6",
}
CHEAPEST_FILL = "FFBBF7D0"
CHEAPEST_FONT_COLOR = "FF14532D"

def norm(s):
    return re.sub(r'\s+', ' ', str(s).strip().lower())

def main():
    if not os.path.exists(CSV_PATH):
        print(f"ERROR: {CSV_PATH} not found. See the instructions at the top of this script.")
        sys.exit(1)
    if not os.path.exists(XLSX_PATH):
        print(f"ERROR: {XLSX_PATH} not found. Put this script in the same folder as your database file.")
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    if "Products" not in wb.sheetnames:
        print("ERROR: no 'Products' sheet found in the workbook.")
        sys.exit(1)
    ws = wb["Products"]

    icon_lookup = {}
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        icon = ws.cell(row=r, column=4).value
        if name and icon:
            icon_lookup[norm(name)] = icon

    rows = []
    with open(CSV_PATH, encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        next(reader)  # header
        for row in reader:
            if row and row[0].strip():
                rows.append(row)

    print(f"Read {len(rows)} products from CSV.")

    max_old_row = ws.max_row
    for r in range(4, max_old_row + 1):
        for c in range(1, 21):
            ws.cell(row=r, column=c).value = None

    today_str = date.today().strftime("%d %B %Y")

    for i, row in enumerate(rows):
        pid, name, cat_raw, coop, morr, ms, wait, ice, alli = row
        r = 4 + i

        m = re.match(r'^(\S+)\s+(.*)$', cat_raw.strip())
        if m and not m.group(1)[0].isalnum():
            icon_from_cat = m.group(1)
            cat_text = m.group(2)
        else:
            icon_from_cat = ''
            cat_text = cat_raw.strip()

        icon = icon_lookup.get(norm(name), icon_from_cat)

        prices = {}
        for label, val in zip(STORE_DISPLAY, [coop, morr, ms, wait, ice, alli]):
            val = val.strip() if val else ''
            prices[label] = float(val) if val else None

        valid_prices = {k: v for k, v in prices.items() if v is not None}
        if valid_prices:
            cheapest_store = min(valid_prices, key=valid_prices.get)
            lowest_price = valid_prices[cheapest_store]
        else:
            cheapest_store = None
            lowest_price = None

        ws.cell(row=r, column=1, value=int(pid))
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=cat_text)
        ws.cell(row=r, column=4, value=icon)
        ws.cell(row=r, column=5, value=prices["CI Co-op"])
        ws.cell(row=r, column=7, value=prices["Morrisons"])
        ws.cell(row=r, column=9, value=prices["M&S Food"])
        ws.cell(row=r, column=11, value=prices["Waitrose"])
        ws.cell(row=r, column=13, value=prices["Iceland"])
        ws.cell(row=r, column=15, value=prices["Alliance"])
        ws.cell(row=r, column=17, value=lowest_price)
        ws.cell(row=r, column=18, value=cheapest_store)
        ws.cell(row=r, column=19, value=today_str)

        # Re-derive each price cell's highlight from today's actual prices,
        # instead of leaving behind whichever store used to be cheapest.
        # Every store within half a penny of the row's lowest price is
        # highlighted (ties are common and all deserve the highlight).
        for store, col in STORE_COLS.items():
            cell = ws.cell(row=r, column=col)
            val = prices[store]
            is_cheapest = val is not None and lowest_price is not None and abs(val - lowest_price) < 0.005
            existing_font = cell.font
            if is_cheapest:
                cell.font = Font(name=existing_font.name, size=existing_font.size, bold=True, color=CHEAPEST_FONT_COLOR)
                cell.fill = PatternFill(fill_type="solid", fgColor=CHEAPEST_FILL)
            else:
                cell.font = Font(name=existing_font.name, size=existing_font.size, bold=False, color=None)
                cell.fill = PatternFill(fill_type="solid", fgColor=NORMAL_FILL[store])

    wb.save(XLSX_PATH)
    print(f"Done. Products sheet rebuilt with {len(rows)} products, dated {today_str}.")
    print("NOTE: open the file in Excel once to let it recalculate the Summary sheet formulas.")

if __name__ == "__main__":
    main()
